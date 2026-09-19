"""txt / docx / pdf からプレーンテキストを抽出する。

出力は「段落を空行 (\\n\\n) で区切ったテキスト」に正規化する。
チャンク分割はこの段落区切りを最優先の意味の切れ目として扱う。
"""
from __future__ import annotations

import io
import re
from pathlib import Path
from typing import BinaryIO

SUPPORTED_EXTENSIONS = {".txt", ".docx", ".pdf"}


def extract_text(source: str | Path | BinaryIO, filename: str | None = None) -> str:
    """ファイルパスまたはバイナリストリームからテキストを抽出する。"""
    if isinstance(source, (str, Path)):
        path = Path(source)
        filename = filename or path.name
        data = path.read_bytes()
    else:
        if filename is None:
            raise ValueError("ストリーム入力の場合は filename が必要です")
        data = source.read()

    ext = Path(filename).suffix.lower()
    if ext == ".txt":
        text = _extract_txt(data)
    elif ext == ".docx":
        text = _extract_docx(data)
    elif ext == ".pdf":
        text = _extract_pdf(data)
    else:
        raise ValueError(f"未対応の拡張子です: {ext}（対応: {', '.join(sorted(SUPPORTED_EXTENSIONS))}）")
    return normalize(text)


def _extract_txt(data: bytes) -> str:
    if data[:2] in (b"\xff\xfe", b"\xfe\xff"):
        return data.decode("utf-16")
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def _extract_docx(data: bytes) -> str:
    import docx  # python-docx

    doc = docx.Document(io.BytesIO(data))
    blocks: list[str] = []
    for para in doc.paragraphs:
        t = para.text.strip()
        if t:
            blocks.append(t)
    # 表のセルも段落として拾う
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                blocks.append(" | ".join(dict.fromkeys(cells)))
    return "\n\n".join(blocks)


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = [(page.extract_text() or "") for page in reader.pages]
    return "\n\n".join(_reflow_pdf_page(p) for p in pages if p.strip())


def _reflow_pdf_page(text: str) -> str:
    """PDF の固定改行を段落に組み直す。

    - 行末ハイフネーション (Über-\\nsetzung) を結合
    - 文末記号で終わる行の後、または空行を段落の切れ目とみなす
    - それ以外の改行はスペースに置換
    """
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)
    lines = [ln.strip() for ln in text.split("\n")]
    paras: list[str] = []
    buf: list[str] = []
    for ln in lines:
        if not ln:
            if buf:
                paras.append(" ".join(buf))
                buf = []
            continue
        buf.append(ln)
        # 短い行 + 文末記号 = 段落末の可能性が高い
        if re.search(r"[.!?:][\"'”»)]?$", ln) and len(ln) < 60:
            paras.append(" ".join(buf))
            buf = []
    if buf:
        paras.append(" ".join(buf))
    return "\n\n".join(paras)


def normalize(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace(" ", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()
