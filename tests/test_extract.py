import io

import docx

from deepl_translator.extract import extract_text


def test_docx_keeps_table_in_document_order():
    doc = docx.Document()
    doc.add_paragraph("Before table.")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Value"
    doc.add_paragraph("After table.")
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    assert extract_text(buf, filename="x.docx") == "Before table.\n\nName | Value\n\nAfter table."
